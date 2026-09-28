from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .analyzer import AnalysisResult
from .constants import REPORT_COLUMNS


def write_csv(results: list[AnalysisResult], path: str | Path) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=REPORT_COLUMNS)
        writer.writeheader()
        writer.writerows(result.report_row() for result in results)
    return output


def write_xlsx(results: list[AnalysisResult], path: str | Path) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Analysis results"
    sheet.append(REPORT_COLUMNS)
    for result in results:
        row = result.report_row()
        sheet.append([row[column] for column in REPORT_COLUMNS])
    header_fill = PatternFill("solid", fgColor="123047")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    widths = (42, 35, 35, 38, 14, 46, 18, 20, 16, 18)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(output)
    return output


def write_manifest(results: list[AnalysisResult], path: str | Path) -> Path:
    output = Path(path)
    payload = {
        "summary": {
            "files": len(results),
            "success": sum(r.processing_status == "Success" for r in results),
            "failure": sum(r.processing_status == "Failure" for r in results),
            "with_violations": sum(r.quality_class == 1 for r in results),
            "manual_review": sum(r.review_priority in {"high", "medium"} for r in results),
        },
        "results": [result.api_dict() for result in results],
    }
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return output
